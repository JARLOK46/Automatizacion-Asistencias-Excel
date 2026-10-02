from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
import hashlib
import os
import shutil
import tempfile
from typing import Iterable
import zipfile
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from .config import AppConfig
from .validation import AttendanceRecordInput
from .db import StoredAttendanceRecord


class WorkbookSafetyError(RuntimeError):
    """Raised when an XLSX candidate cannot be proven safe to replace."""


@dataclass(frozen=True)
class WorkbookCandidate:
    source: Path
    output: Path
    backup: Path | None
    rows_written: tuple[int, ...]


def locate_workbook(config: AppConfig) -> Path:
    path = config.workbook_path
    if not path.is_file():
        raise WorkbookSafetyError(f"configured workbook does not exist: {path}")
    return path


def create_backup(workbook: Path, backups_dir: Path) -> Path:
    if not workbook.is_file():
        raise WorkbookSafetyError(f"cannot back up missing workbook: {workbook}")
    backups_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    backup = backups_dir / f"{workbook.stem}.{stamp}{workbook.suffix}"
    shutil.copy2(workbook, backup)
    return backup


def load_workbook_copy(workbook: Path, directory: Path | None = None):
    if not workbook.is_file():
        raise WorkbookSafetyError(f"workbook does not exist: {workbook}")
    directory = directory or Path(tempfile.mkdtemp(prefix="attendance-workbook-"))
    directory.mkdir(parents=True, exist_ok=True)
    copied = directory / workbook.name
    shutil.copy2(workbook, copied)
    return copied, load_workbook(copied)


def _nonempty(value: object) -> bool:
    return value is not None and value != ""


def find_attendance_end(sheet, start_row: int = 17) -> int:
    """Find the last populated attendance row, ignoring template sequence numbers in A."""
    rows = {cell.row for cell in sheet._cells.values() if cell.row >= start_row and 2 <= cell.column <= 14 and _nonempty(cell.value)}
    return max(rows, default=start_row - 1)


def clear_attendance_values(sheet, start_row: int = 17) -> None:
    for cell in sheet._cells.values():
        if cell.row >= start_row and 1 <= cell.column <= 14:
            cell.value = None


def record_values(record: AttendanceRecordInput, index: int) -> tuple[object, ...]:
    return (index, record.full_name, record.document_type, record.document_number, record.gender, record.training_sheet_number, record.modality, record.program, record.schedule, record.education_level, record.training_center, record.email, record.phone, record.registered_at_local)


def append_records(sheet, records: Iterable[AttendanceRecordInput]) -> tuple[int, ...]:
    row = find_attendance_end(sheet) + 1
    written: list[int] = []
    for record in records:
        sequence = int(sheet.cell(row=row, column=1).value or row - 16)
        for column, value in enumerate(record_values(record, sequence), start=1):
            sheet.cell(row=row, column=column).value = value
        written.append(row)
        row += 1
    return tuple(written)


def _cell_signature(sheet, row: int, column: int) -> tuple[object, int]:
    cell = sheet.cell(row=row, column=column)
    return cell.value, cell.style_id


def _zip_signature(path: Path) -> dict[str, str]:
    with ZipFile(path) as archive:
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if not name.startswith("xl/worksheets/sheet")}


def _has_drawings(path: Path) -> bool:
    with ZipFile(path) as archive:
        return any(name.startswith("xl/drawings/") and name.endswith(".xml") for name in archive.namelist())


def protected_signature(workbook_path: Path) -> tuple[object, ...]:
    if workbook_path.suffix not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        import io
        workbook = load_workbook(io.BytesIO(workbook_path.read_bytes()), read_only=False)
    else:
        workbook = load_workbook(workbook_path, read_only=False)
    sheets = []
    for name in workbook.sheetnames:
        sheet = workbook[name]
        protected_cells = tuple(_cell_signature(sheet, row, column) for row in range(1, 17) for column in range(1, 15))
        sheets.append((name, tuple(str(r) for r in sheet.merged_cells.ranges), sheet.max_column, protected_cells))
    return tuple(sheets)


def _xml_text(value: object) -> str:
    text = "" if value is None else str(value)
    # XML 1.0 permits tab, LF, and CR, but not the other C0 controls.
    if any(ord(character) < 0x20 and ord(character) not in {0x09, 0x0A, 0x0D} for character in text):
        raise WorkbookSafetyError("attendance value contains an XML 1.0-invalid control character")
    return text


def _set_cell(cell, value: object) -> None:
    # Cell type metadata is part of the value representation.  Leaving a stale
    # type behind can make Excel reject an otherwise well-formed worksheet.
    cell.attrib.pop("t", None)
    for child in list(cell):
        if child.tag.rsplit("}", 1)[-1] in {"v", "is", "f"}:
            cell.remove(child)
    if value is None or value == "":
        return
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    if isinstance(value, int):
        child = ET.SubElement(cell, f"{{{ns}}}v")
        child.text = str(value)
    else:
        cell.set("t", "inlineStr")
        inline = ET.SubElement(cell, f"{{{ns}}}is")
        text = ET.SubElement(inline, f"{{{ns}}}t")
        text.text = _xml_text(value)


def _sort_cells(row, spreadsheet_ns: str) -> None:
    """Keep row cell elements in the canonical left-to-right order."""
    cells = row.findall(f"{{{spreadsheet_ns}}}c")
    ordered = sorted(cells, key=lambda cell: _column_number("".join(c for c in cell.attrib.get("r", "") if c.isalpha())))
    if cells == ordered:
        return
    positions = [index for index, child in enumerate(list(row)) if child in cells]
    children = list(row)
    for index, cell in zip(positions, ordered):
        children[index] = cell
    row[:] = children


def _xml_project(source: Path, output: Path, records: Iterable[AttendanceRecordInput], reset: bool = False) -> tuple[int, ...]:
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(source) as archive:
        entries = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    part = _resolve_sheet_part(source, "Formato")
    root = ET.fromstring(entries[part])
    sheet_data = root.find(f"{{{spreadsheet_ns}}}sheetData")
    if sheet_data is None:
        raise WorkbookSafetyError("Formato sheet has no sheetData")
    rows = {int(row.attrib["r"]): row for row in sheet_data.findall(f"{{{spreadsheet_ns}}}row")}
    if reset:
        for row in rows.values():
            if int(row.attrib["r"]) >= 17:
                for cell in row.findall(f"{{{spreadsheet_ns}}}c"):
                    if 1 <= _column_number("".join(c for c in cell.attrib.get("r", "") if c.isalpha())) <= 14:
                        _set_cell(cell, None)
    end = 16
    for number, row in rows.items():
        if number < 17:
            continue
        for cell in row.findall(f"{{{spreadsheet_ns}}}c"):
            col = "".join(c for c in cell.attrib.get("r", "") if c.isalpha())
            if 2 <= _column_number(col) <= 14 and any(child.tag.rsplit("}", 1)[-1] in {"v", "is"} for child in cell):
                end = max(end, number)
    written = []
    for offset, record in enumerate(records):
        number = end + offset + 1
        row = rows.get(number)
        if row is None:
            row = ET.Element(f"{{{spreadsheet_ns}}}row", {"r": str(number), "spans": "1:14"})
            sheet_data.append(row)
        elif row.attrib.get("spans"):
            row.set("spans", "1:14")
        cells = {"".join(c for c in cell.attrib.get("r", "") if c.isalpha()): cell for cell in row.findall(f"{{{spreadsheet_ns}}}c")}
        for column, value in enumerate(record_values(record, offset + 1), 1):
            letter = ""
            n = column
            while n:
                n, rem = divmod(n - 1, 26); letter = chr(65 + rem) + letter
            cell = cells.get(letter)
            if cell is None:
                cell = ET.SubElement(row, f"{{{spreadsheet_ns}}}c", {"r": f"{letter}{number}"})
            _set_cell(cell, value)
        _sort_cells(row, spreadsheet_ns)
        written.append(number)
    entries[part] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(output, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content, compress_type=zipfile.ZIP_DEFLATED)
    return tuple(written)


def verify_candidate(source: Path, candidate: Path) -> None:
    if protected_signature(source) != protected_signature(candidate):
        raise WorkbookSafetyError("protected rows, merges, sheet names, or styles changed")
    if _zip_signature(source) != _zip_signature(candidate):
        raise WorkbookSafetyError("non-worksheet workbook structure changed")


class WorkbookProjector:
    def __init__(self, config: AppConfig):
        self.config = config

    def project(self, records: Iterable[AttendanceRecordInput], reset: bool = False) -> WorkbookCandidate:
        source = locate_workbook(self.config)
        backup = create_backup(source, self.config.backups_path) if reset else None
        with tempfile.TemporaryDirectory(prefix="attendance-export-") as directory:
            output = Path(directory) / source.name
            shutil.copy2(source, output)
            rows = _xml_project(source, output, records, reset=reset)
            verify_candidate(source, output)
            replacement = source.with_suffix(source.suffix + ".pending-replace")
            shutil.copy2(output, replacement)
            verify_candidate(source, replacement)
            os.replace(replacement, source)
            return WorkbookCandidate(source, output, backup, rows)


def _resolve_sheet_part(workbook_path: Path, sheet_name: str) -> str:
    ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "office": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    with ZipFile(workbook_path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in relationships}
        for sheet in workbook.findall("main:sheets/main:sheet", ns):
            if sheet.attrib.get("name") == sheet_name:
                target = rel_targets[sheet.attrib["{" + ns['office'] + "}id"]]
                return str((Path("xl") / target).as_posix()).replace("xl/../", "")
    raise WorkbookSafetyError(f"workbook has no {sheet_name} sheet")


def _column_number(column: str) -> int:
    number = 0
    for character in column.upper():
        number = number * 26 + ord(character) - ord("A") + 1
    return number


def reset_workbook_values(workbook: Path, backups_dir: Path) -> Path:
    """Clear attendance values while preserving package parts and cell styles."""
    source = workbook.resolve()
    backup = create_backup(source, backups_dir)
    sheet_part = _resolve_sheet_part(source, "Formato")
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(source) as archive:
        entries = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    root = ET.fromstring(entries[sheet_part])
    for row in root.findall(f"{{{spreadsheet_ns}}}sheetData/{{{spreadsheet_ns}}}row"):
        if int(row.attrib.get("r", "0")) < 17:
            continue
        for cell in row.findall(f"{{{spreadsheet_ns}}}c"):
            column = "".join(character for character in cell.attrib.get("r", "") if character.isalpha())
            if column and 1 <= _column_number(column) <= 14:
                cell.attrib.pop("t", None)
                for child in list(cell):
                    if child.tag.rsplit("}", 1)[-1] in {"v", "is", "f"}:
                        cell.remove(child)
    entries[sheet_part] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    replacement = source.with_suffix(source.suffix + ".pending-reset")
    with ZipFile(replacement, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    os.replace(replacement, source)
    return backup
